from datetime import date
from io import BytesIO

from flask import Blueprint, render_template, request, send_file
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment

# >>> Troque pela sua função/classe de conexão com o MySQL <<<
from core.database import conectar_banco   # deve retornar uma conexão mysql.connector

relatorios_bp = Blueprint("relatorios", __name__)

TIPOS = {
    "estoque":      "Estoque atual",
    "movimentacao": "Entradas e saídas",
    "baixo":        "Estoque baixo e crítico",
    "validade":     "Validade dos medicamentos",
    "financeiro":   "Resumo financeiro",
}


# ---------------------------------------------------------------- helpers
def consultar(sql, params=()):
    conn = conectar_banco.connect()
    cur = conn.cursor(dictionary=True)
    cur.execute(sql, params)
    linhas = cur.fetchall()
    cur.close()
    conn.close()
    return linhas


@relatorios_bp.app_template_filter("brl")
def brl(valor):
    valor = float(valor or 0)
    return "R$ " + f"{valor:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")


@relatorios_bp.app_template_filter("inteiro")
def inteiro(valor):
    return f"{int(valor or 0):,}".replace(",", ".")


def ler_filtros():
    hoje = date.today()
    return {
        "tipo": request.args.get("tipo", "estoque"),
        "data_inicial": request.args.get("data_inicial", hoje.replace(day=1).isoformat()),
        "data_final": request.args.get("data_final", hoje.isoformat()),
        "fornecedor_id": request.args.get("fornecedor_id", ""),
        "categoria": request.args.get("categoria", ""),
        "tipo_mov": request.args.get("tipo_mov", ""),
        "prazo": request.args.get("prazo", "90"),
    }


# ---------------------------------------------------------------- relatórios
# Cada função devolve (colunas, linhas, resumo)
# colunas = [(chave, rótulo, formato)]  formato: texto|int|moeda|data|datahora|pct|badge
# resumo  = [(rótulo, valor, formato, ícone)]

def rel_estoque(f):
    where, params = ["p.ativo = TRUE"], []
    if f["fornecedor_id"]:
        where.append("p.fornecedor_id = %s"); params.append(f["fornecedor_id"])
    if f["categoria"]:
        where.append("p.categoria = %s"); params.append(f["categoria"])

    linhas = consultar(f"""
        SELECT p.id, p.nome, p.categoria, fo.nome_fornecedor,
               p.quantidade_estoque, p.estoque_minimo,
               p.preco_custo, p.preco_venda,
               p.quantidade_estoque * p.preco_custo AS valor_custo,
               p.quantidade_estoque * p.preco_venda AS valor_venda
        FROM produto p
        JOIN fornecedor fo ON fo.id = p.fornecedor_id
        WHERE {' AND '.join(where)}
        ORDER BY p.nome
    """, params)

    colunas = [
        ("nome", "Produto", "texto"), ("categoria", "Categoria", "texto"),
        ("nome_fornecedor", "Fornecedor", "texto"),
        ("quantidade_estoque", "Qtd.", "int"), ("estoque_minimo", "Mínimo", "int"),
        ("preco_custo", "Custo un.", "moeda"), ("preco_venda", "Venda un.", "moeda"),
        ("valor_custo", "Total custo", "moeda"), ("valor_venda", "Total venda", "moeda"),
    ]
    resumo = [
        ("Produtos", len(linhas), "int", "bi-capsule"),
        ("Unidades em estoque", sum(l["quantidade_estoque"] for l in linhas), "int", "bi-box-seam"),
        ("Valor a custo", sum(l["valor_custo"] for l in linhas), "moeda", "bi-cash-stack"),
        ("Valor de venda", sum(l["valor_venda"] for l in linhas), "moeda", "bi-graph-up-arrow"),
    ]
    return colunas, linhas, resumo


def rel_movimentacao(f):
    where, params = ["DATE(m.data_mov) BETWEEN %s AND %s"], [f["data_inicial"], f["data_final"]]
    if f["tipo_mov"] in ("ENTRADA", "SAIDA"):
        where.append("m.tipo = %s"); params.append(f["tipo_mov"])
    if f["fornecedor_id"]:
        where.append("m.fornecedor_id = %s"); params.append(f["fornecedor_id"])

    linhas = consultar(f"""
        SELECT m.id, m.data_mov, m.tipo, p.nome AS produto, m.quantidade, m.valor_total,
               COALESCE(fo.nome_fornecedor, c.nome, '—') AS origem_destino,
               COALESCE(e.observacao, s.observacao, '—') AS motivo
        FROM movimentacao m
        JOIN produto p                 ON p.id  = m.produto_id
        LEFT JOIN fornecedor fo        ON fo.id = m.fornecedor_id
        LEFT JOIN clientes_cadastro c  ON c.id  = m.cliente_id
        LEFT JOIN entrada e            ON e.id  = m.entrada_id
        LEFT JOIN saida s              ON s.id  = m.saida_id
        WHERE {' AND '.join(where)}
        ORDER BY m.data_mov DESC
    """, params)

    colunas = [
        ("data_mov", "Data", "datahora"), ("tipo", "Tipo", "badge"),
        ("produto", "Produto", "texto"), ("quantidade", "Qtd.", "int"),
        ("valor_total", "Valor", "moeda"),
        ("origem_destino", "Fornecedor / Cliente", "texto"), ("motivo", "Observação", "texto"),
    ]
    ent = [l for l in linhas if l["tipo"] == "ENTRADA"]
    sai = [l for l in linhas if l["tipo"] == "SAIDA"]
    resumo = [
        ("Movimentações", len(linhas), "int", "bi-arrow-left-right"),
        ("Unidades que entraram", sum(l["quantidade"] for l in ent), "int", "bi-box-arrow-in-down"),
        ("Unidades que saíram", sum(l["quantidade"] for l in sai), "int", "bi-box-arrow-up"),
        ("Saldo (entradas − saídas)",
         sum(l["quantidade"] for l in ent) - sum(l["quantidade"] for l in sai), "int", "bi-calculator"),
    ]
    return colunas, linhas, resumo


def rel_baixo(f):
    where, params = ["p.ativo = TRUE", "p.quantidade_estoque <= p.estoque_minimo"], []
    if f["fornecedor_id"]:
        where.append("p.fornecedor_id = %s"); params.append(f["fornecedor_id"])
    if f["categoria"]:
        where.append("p.categoria = %s"); params.append(f["categoria"])

    linhas = consultar(f"""
        SELECT p.id, p.nome, p.categoria, fo.nome_fornecedor,
               p.quantidade_estoque, p.estoque_minimo,
               CASE WHEN p.quantidade_estoque = 0
                      OR p.quantidade_estoque <= p.estoque_minimo * 0.5
                    THEN 'CRITICO' ELSE 'BAIXO' END AS situacao,
               GREATEST(p.estoque_minimo * 2 - p.quantidade_estoque, 0) AS sugestao_reposicao,
               GREATEST(p.estoque_minimo * 2 - p.quantidade_estoque, 0) * p.preco_custo AS custo_reposicao
        FROM produto p
        JOIN fornecedor fo ON fo.id = p.fornecedor_id
        WHERE {' AND '.join(where)}
        ORDER BY (p.quantidade_estoque - p.estoque_minimo), p.nome
    """, params)

    colunas = [
        ("nome", "Produto", "texto"), ("categoria", "Categoria", "texto"),
        ("nome_fornecedor", "Fornecedor", "texto"),
        ("quantidade_estoque", "Qtd.", "int"), ("estoque_minimo", "Mínimo", "int"),
        ("situacao", "Situação", "badge"),
        ("sugestao_reposicao", "Repor (un.)", "int"), ("custo_reposicao", "Custo reposição", "moeda"),
    ]
    resumo = [
        ("Críticos", sum(1 for l in linhas if l["situacao"] == "CRITICO"), "int", "bi-exclamation-octagon"),
        ("Baixos", sum(1 for l in linhas if l["situacao"] == "BAIXO"), "int", "bi-exclamation-triangle"),
        ("Unidades a repor", sum(l["sugestao_reposicao"] for l in linhas), "int", "bi-box-seam"),
        ("Custo estimado", sum(l["custo_reposicao"] for l in linhas), "moeda", "bi-cash-stack"),
    ]
    return colunas, linhas, resumo


def rel_validade(f):
    where, params = ["p.ativo = TRUE", "p.data_validade IS NOT NULL"], []
    if f["prazo"] == "vencidos":
        where.append("p.data_validade < CURDATE()")
    else:
        where.append("p.data_validade <= DATE_ADD(CURDATE(), INTERVAL %s DAY)")
        params.append(int(f["prazo"]))
    if f["fornecedor_id"]:
        where.append("p.fornecedor_id = %s"); params.append(f["fornecedor_id"])
    if f["categoria"]:
        where.append("p.categoria = %s"); params.append(f["categoria"])

    linhas = consultar(f"""
        SELECT p.id, p.nome, p.categoria, fo.nome_fornecedor, p.quantidade_estoque,
               p.data_validade, DATEDIFF(p.data_validade, CURDATE()) AS dias_restantes,
               CASE WHEN p.data_validade < CURDATE() THEN 'VENCIDO'
                    WHEN p.data_validade <= DATE_ADD(CURDATE(), INTERVAL 30 DAY) THEN 'VENCE EM 30 DIAS'
                    ELSE 'A VENCER' END AS situacao,
               p.quantidade_estoque * p.preco_custo AS valor_em_risco
        FROM produto p
        JOIN fornecedor fo ON fo.id = p.fornecedor_id
        WHERE {' AND '.join(where)}
        ORDER BY p.data_validade, p.nome
    """, params)

    colunas = [
        ("nome", "Produto", "texto"), ("categoria", "Categoria", "texto"),
        ("nome_fornecedor", "Fornecedor", "texto"), ("quantidade_estoque", "Qtd.", "int"),
        ("data_validade", "Validade", "data"), ("dias_restantes", "Dias", "int"),
        ("situacao", "Situação", "badge"), ("valor_em_risco", "Valor em risco", "moeda"),
    ]
    resumo = [
        ("Vencidos", sum(1 for l in linhas if l["situacao"] == "VENCIDO"), "int", "bi-x-octagon"),
        ("Vencem em 30 dias", sum(1 for l in linhas if l["situacao"] == "VENCE EM 30 DIAS"), "int", "bi-hourglass-split"),
        ("Unidades afetadas", sum(l["quantidade_estoque"] for l in linhas), "int", "bi-box-seam"),
        ("Valor em risco", sum(l["valor_em_risco"] for l in linhas), "moeda", "bi-cash-stack"),
    ]
    return colunas, linhas, resumo


def rel_financeiro(f):
    linhas = consultar("""
        SELECT p.categoria,
               COUNT(*)                                        AS produtos,
               SUM(p.quantidade_estoque)                       AS unidades,
               SUM(p.quantidade_estoque * p.preco_custo)       AS valor_custo,
               SUM(p.quantidade_estoque * p.preco_venda)       AS valor_venda,
               SUM(p.quantidade_estoque * (p.preco_venda - p.preco_custo)) AS lucro_potencial,
               ROUND((SUM(p.quantidade_estoque * p.preco_venda) - SUM(p.quantidade_estoque * p.preco_custo))
                     / NULLIF(SUM(p.quantidade_estoque * p.preco_venda), 0) * 100, 1) AS margem
        FROM produto p
        WHERE p.ativo = TRUE
        GROUP BY p.categoria
        ORDER BY valor_custo DESC
    """)

    periodo = consultar("""
        SELECT tipo, COALESCE(SUM(valor_total), 0) AS total
        FROM movimentacao
        WHERE DATE(data_mov) BETWEEN %s AND %s
        GROUP BY tipo
    """, (f["data_inicial"], f["data_final"]))
    mov = {r["tipo"]: r["total"] for r in periodo}

    custo = sum(l["valor_custo"] for l in linhas)
    venda = sum(l["valor_venda"] for l in linhas)
    margem = round((venda - custo) / venda * 100, 1) if venda else 0

    colunas = [
        ("categoria", "Categoria", "texto"), ("produtos", "Produtos", "int"),
        ("unidades", "Unidades", "int"), ("valor_custo", "Valor a custo", "moeda"),
        ("valor_venda", "Valor de venda", "moeda"),
        ("lucro_potencial", "Lucro potencial", "moeda"), ("margem", "Margem %", "pct"),
    ]
    resumo = [
        ("Valor do estoque (custo)", custo, "moeda", "bi-cash-stack"),
        ("Valor potencial de venda", venda, "moeda", "bi-graph-up-arrow"),
        ("Lucro potencial", venda - custo, "moeda", "bi-piggy-bank"),
        ("Margem estimada", margem, "pct", "bi-percent"),
        ("Entradas no período", mov.get("ENTRADA", 0), "moeda", "bi-box-arrow-in-down"),
        ("Saídas no período", mov.get("SAIDA", 0), "moeda", "bi-box-arrow-up"),
    ]
    return colunas, linhas, resumo


GERADORES = {
    "estoque": rel_estoque, "movimentacao": rel_movimentacao, "baixo": rel_baixo,
    "validade": rel_validade, "financeiro": rel_financeiro,
}


# ---------------------------------------------------------------- rotas
@relatorios_bp.route("/relatorios")
def tela_relatorios():
    f = ler_filtros()
    if f["tipo"] not in GERADORES:
        f["tipo"] = "estoque"
    colunas, linhas, resumo = GERADORES[f["tipo"]](f)

    fornecedores = consultar("SELECT id, nome_fornecedor FROM fornecedor WHERE ativo = TRUE ORDER BY nome_fornecedor")
    categorias = [r["categoria"] for r in consultar(
        "SELECT DISTINCT categoria FROM produto WHERE ativo = TRUE ORDER BY categoria")]

    return render_template(
        "tela_relatorios.html",
        f=f, tipos=TIPOS, colunas=colunas, linhas=linhas, resumo=resumo,
        fornecedores=fornecedores, categorias=categorias,
        titulo_relatorio=TIPOS[f["tipo"]], gerado_em=date.today().strftime("%d/%m/%Y"),
    )


@relatorios_bp.route("/relatorios/exportar")
def exportar_excel():
    f = ler_filtros()
    if f["tipo"] not in GERADORES:
        f["tipo"] = "estoque"
    colunas, linhas, resumo = GERADORES[f["tipo"]](f)

    wb = Workbook()
    ws = wb.active
    ws.title = TIPOS[f["tipo"]][:31]

    ws.append([f"MEDSTOCK — {TIPOS[f['tipo']]}"])
    ws["A1"].font = Font(bold=True, size=14, color="0B3A7A")
    ws.append([f"Gerado em {date.today().strftime('%d/%m/%Y')}"])
    ws.append([])

    ws.append([c[1] for c in colunas])
    cab = ws.max_row
    for cel in ws[cab]:
        cel.font = Font(bold=True, color="FFFFFF")
        cel.fill = PatternFill("solid", fgColor="0E8FA3")
        cel.alignment = Alignment(horizontal="center")

    for l in linhas:
        ws.append([l[c[0]] for c in colunas])

    for i, (_, _, fmt) in enumerate(colunas, start=1):
        letra = ws.cell(row=cab, column=i).column_letter
        ws.column_dimensions[letra].width = 22
        for r in range(cab + 1, ws.max_row + 1):
            if fmt == "moeda":
                ws.cell(row=r, column=i).number_format = '"R$" #,##0.00'
            elif fmt == "data":
                ws.cell(row=r, column=i).number_format = "DD/MM/YYYY"
            elif fmt == "datahora":
                ws.cell(row=r, column=i).number_format = "DD/MM/YYYY HH:MM"

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return send_file(
        buf, as_attachment=True,
        download_name=f"relatorio_{f['tipo']}_{date.today().isoformat()}.xlsx",
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )